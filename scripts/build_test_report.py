#!/usr/bin/env python3
"""Build Word docs from manual screenshots + API test report."""
from __future__ import annotations

import json
import re
import subprocess
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Inches

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "docs" / "manual_screenshots"
DOCS = ROOT / "docs" / "screen_docs"
API_DIR = ROOT.parent / "aneuso_api_nodeexpressjs"


def pretty_title(name: str) -> str:
    name = Path(name).stem
    name = re.sub(r"^\d+_", "", name)
    return name.replace("_", " ").title()


def run_api_smoke() -> str:
    script = API_DIR / "scripts" / "smoke-test-api.js"
    if not script.exists():
        return "API smoke script not found."
    try:
        r = subprocess.run(
            ["node", str(script)],
            cwd=str(API_DIR),
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=30,
        )
        return (r.stdout or "") + (r.stderr or "")
    except Exception as e:
        return f"API smoke test could not run: {e}"


def main() -> None:
    DOCS.mkdir(parents=True, exist_ok=True)
    SHOTS.mkdir(parents=True, exist_ok=True)

    pngs = sorted(SHOTS.glob("*.png"))
    api_log = run_api_smoke()

    summary = Document()
    summary.add_heading("ANEUSO Application Test Report", 0)
    summary.add_paragraph(f"Generated: {datetime.now():%Y-%m-%d %H:%M}")
    summary.add_heading("API smoke test", level=1)
    summary.add_paragraph(api_log or "No output")

    summary.add_heading("Screens captured", level=1)
    if not pngs:
        summary.add_paragraph(
            "No screenshots yet. Run the app (flutter run -d windows), open each screen, then:"
        )
        summary.add_paragraph(
            'powershell -File scripts/capture_aneuso_window.ps1 -Name "01_login"'
        )
    else:
        for png in pngs:
            title = pretty_title(png.name)
            doc = Document()
            doc.add_heading(f"ANEUSO — {title}", 0)
            doc.add_paragraph(f"File: {png.name}")
            doc.add_picture(str(png), width=Inches(3.8))
            out = DOCS / f"{png.stem}.docx"
            doc.save(out)
            summary.add_paragraph(f"{title} → {out.name}")

    summary.save(ROOT / "docs" / "ANEUSO_Test_Report.docx")
    print(f"Report: {ROOT / 'docs' / 'ANEUSO_Test_Report.docx'}")
    print(f"Screen docs: {DOCS} ({len(pngs)} images)")


if __name__ == "__main__":
    main()
