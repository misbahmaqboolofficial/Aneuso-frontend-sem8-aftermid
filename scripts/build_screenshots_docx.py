#!/usr/bin/env python3
"""Build single screenshots.docx with captioned screen images."""
from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "test" / "screenshots"
OUT = ROOT / "docs" / "screenshots.docx"


def caption(name: str) -> str:
    stem = Path(name).stem
    stem = re.sub(r"^\d+_", "", stem)
    return stem.replace("_", " ").strip().title()


def main() -> None:
    pngs = sorted(p for p in SHOTS.glob("*.png") if p.stat().st_size > 1000)
    if not pngs:
        raise SystemExit(f"No PNG files in {SHOTS}. Run: flutter test test/screen_capture_test.dart")

    doc = Document()
    title = doc.add_heading("ANEUSO Application Screenshots", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"Generated: {datetime.now():%Y-%m-%d %H:%M}")
    doc.add_paragraph(
        "Screens captured from the Flutter app UI (390×844 phone layout). "
        "Each image shows one screen with its name as caption."
    )

    for i, png in enumerate(pngs, start=1):
        label = caption(png.name)
        doc.add_heading(f"{i}. {label}", level=1)
        p = doc.add_paragraph(label)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.bold = True
            run.font.size = Pt(14)
        doc.add_picture(str(png), width=Inches(3.4))
        doc.add_paragraph()

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Created {OUT} with {len(pngs)} screens")


if __name__ == "__main__":
    main()
