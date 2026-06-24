#!/usr/bin/env python3
"""Create one Word document per screenshot PNG."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SCREENSHOTS = ROOT / "test" / "screenshots"
OUTPUT = ROOT / "docs" / "screen_docs"
SUMMARY = ROOT / "docs" / "ANEUSO_Screen_Test_Summary.docx"


def pretty_title(filename: str) -> str:
    name = Path(filename).stem
    name = re.sub(r"^\d+_", "", name)
    return name.replace("_", " ").title()


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    pngs = sorted(SCREENSHOTS.glob("*.png"))
    if not pngs:
        raise SystemExit(f"No PNG files in {SCREENSHOTS}. Run flutter test first.")

    summary = Document()
    summary.add_heading("ANEUSO Application Screen Test Summary", 0)
    summary.add_paragraph(
        "Automated widget-test screenshots of main application screens. "
        "Each screen has a separate Word document in the screen_docs folder."
    )

    for png in pngs:
        title = pretty_title(png.name)
        doc = Document()
        doc.add_heading(f"ANEUSO — {title}", 0)
        doc.add_paragraph(f"Screen file: {png.name}")
        doc.add_paragraph(
            "Captured via Flutter widget golden test (390×844 logical pixels)."
        )
        doc.add_picture(str(png), width=Inches(3.6))
        out_path = OUTPUT / f"{png.stem}.docx"
        doc.save(out_path)

        summary.add_heading(title, level=2)
        summary.add_paragraph(f"Document: {out_path.name}")

    summary.save(SUMMARY)
    print(f"Created {len(pngs)} screen documents in {OUTPUT}")
    print(f"Summary: {SUMMARY}")


if __name__ == "__main__":
    main()
